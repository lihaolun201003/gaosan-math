# -*- coding: utf-8 -*-
"""Build 复数专题_课堂板书版.pptx : 9 slides, native editable Chinese + XeLaTeX vector math."""
import json, os, copy
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from lxml import etree

BASE = os.path.dirname(os.path.abspath(__file__))
OUT  = os.path.dirname(BASE)                      # deliver into session working dir
PT   = 12700
SP   = 65536.0
EMU_PER_SP = PT / SP
FONT = "\u5fae\u8f6f\u96c5\u9ed1"                 # 微软雅黑
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
RNS     = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS    = "http://schemas.openxmlformats.org/drawingml/2006/main"

C = json.load(open(os.path.join(BASE, "content.json"), encoding="utf-8"))
MAN = json.load(open(os.path.join(BASE, "manifest.json"), encoding="utf-8"))
L = C["layout"]
SW, SH = C["slide_pt"]["w"], C["slide_pt"]["h"]
K = float(L["body_size"]) / float(L["math_base_pt"])
A_T, D_T = 0.80 * L["body_size"], 0.14 * L["body_size"]

def emu(pt): return int(round(pt * PT))
def sp2emu(sp, k): return int(round(sp * k * EMU_PER_SP))

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(emu(SW)), Emu(emu(SH))
BLANK = prs.slide_layouts[6]

def new_slide():
    s = prs.slides.add_slide(BLANK)
    # explicit pure-white background
    cSld = s._element.find(qn('p:cSld'))
    bg = etree.SubElement(cSld, qn('p:bg'))
    cSld.remove(bg); cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, qn('p:bgPr'))
    fill = etree.SubElement(bgPr, qn('a:solidFill'))
    clr = etree.SubElement(fill, qn('a:srgbClr')); clr.set('val', 'FFFFFF')
    etree.SubElement(bgPr, qn('a:effectLst'))
    return s

def style_run(run, size, bold, color):
    f = run.font
    f.size = Emu(emu(size)) if False else __import__('pptx.util', fromlist=['Pt']).Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        rPr.append(rPr.makeelement(qn(tag), {"typeface": FONT}))

def add_text(slide, text, left_pt, baseline_pt, size, color, bold=False,
             align=PP_ALIGN.LEFT, width_pt=None):
    """Top-anchored, zero-inset textbox.  Empirical rule (calibrated for 微软雅黑):
       first-line baseline == box_top + font_size."""
    w = width_pt if width_pt is not None else (len(text) * size + 6)
    top = baseline_pt - size
    tb = slide.shapes.add_textbox(Emu(emu(left_pt)), Emu(emu(top)), Emu(emu(w)), Emu(emu(size * 2)))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    style_run(r, size, bold, color)
    return tb

def add_svg_pic(slide, svg_path, png_path, left_pt, top_pt, w_pt, h_pt):
    """Insert picture whose primary format is SVG (vector) with a PNG fallback."""
    pic = slide.shapes.add_picture(png_path, Emu(emu(left_pt)), Emu(emu(top_pt)),
                                   Emu(emu(w_pt)), Emu(emu(h_pt)))
    pkg = slide.part.package
    partname = pkg.next_partname(PackURI('/ppt/media/image%d.svg'))
    svg_part = Part(partname, 'image/svg+xml', pkg, open(svg_path, 'rb').read())
    rId = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = pic._element.blipFill.blip
    extLst = etree.SubElement(blip, qn('a:extLst'))
    ext = etree.SubElement(extLst, qn('a:ext'))
    ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    sb = etree.SubElement(ext, '{%s}svgBlip' % NS_ASVG)
    sb.set('{%s}embed' % RNS, rId)
    pic._element.nvPicPr.cNvPr.set('descr', 'LaTeX formula (SVG)')
    return pic

def add_math(slide, name, left_pt, baseline_pt):
    m = MAN[name]
    w = m["wd_sp"] / SP * K
    ht = m["ht_sp"] / SP * K
    dp = m["dp_sp"] / SP * K
    add_svg_pic(slide, os.path.join(BASE, m["svg"]), os.path.join(BASE, m["png"]),
                left_pt, baseline_pt - ht, w, ht + dp)
    return w

# ---------------- cover ----------------
s = new_slide()
add_text(s, C["cover"]["text"], 0, L["cover_baseline"], L["cover_size"],
         L["title_color"], bold=True, align=PP_ALIGN.CENTER, width_pt=SW)

# ---------------- question slides ----------------
audit = []
for sl in C["slides"]:
    s = new_slide()
    add_text(s, sl["title"], L["margin_l"], L["title_box_top"] + L["title_size"],
             L["title_size"], L["title_color"], bold=True)
    y = float(L["content_top"])
    for ri, row in enumerate(sl["rows"]):
        A = D = 0.0
        for seg in row:
            if seg["k"] == "m":
                m = MAN[seg["s"]]
                A = max(A, m["ht_sp"] / SP * K); D = max(D, m["dp_sp"] / SP * K)
            else:
                A = max(A, A_T); D = max(D, D_T)
        if ri: y += L["row_gap"]
        base = y + A
        display_only = (len(row) == 1 and row[0].get("display"))
        x = L["margin_l"] + (L["display_indent"] if display_only else 0)
        prev_label = False
        for si, seg in enumerate(row):
            if si: x += L["label_gap"] if prev_label else L["seg_gap"]
            if seg["k"] == "m":
                x += add_math(s, seg["s"], x, base)
            else:
                is_label = (seg["k"] == "label")
                add_text(s, seg["s"], x, base, L["body_size"],
                         L["label_color"] if is_label else L["body_color"], bold=is_label)
                x += len(seg["s"]) * L["body_size"]
            prev_label = (seg["k"] == "label")
        y = base + D
    audit.append((sl["n"], round(y, 1), round((540 - y) / 540 * 100, 1)))

out = os.path.join(OUT, "\u590d\u6570\u4e13\u9898_\u8bfe\u5802\u677f\u4e66\u7248.pptx")
prs.save(out)
print("saved:", out, os.path.getsize(out), "bytes, slides =", len(prs.slides.__iter__.__self__._sldIdLst))
for n, b, blank in audit:
    print("  slide %d: contentBottom=%.1fpt  blank=%.1f%%" % (n, b, blank))
