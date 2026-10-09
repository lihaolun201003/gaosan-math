# -*- coding: utf-8 -*-
"""Build 充分必要条件_9大题型课堂板书.pptx

1 cover + 9 question slides, 16:9 (960x540 pt), pure-white background,
native 微软雅黑 text for Chinese, XeLaTeX-compiled vector SVG for every math
formula, and at least 60% of each slide left completely blank for handwriting.
"""
import json, os, sys

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(HERE)                      # deliver into the session working dir
sys.path.insert(0, HERE)
import layout
import metrics

PT = 12700
FONT = "\u5fae\u8f6f\u96c5\u9ed1"                # 微软雅黑
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
OUTNAME = "\u5145\u5206\u5fc5\u8981\u6761\u4ef6_9\u5927\u9898\u578b\u8bfe\u5802\u677f\u4e66.pptx"


def emu(pt):
    return int(round(pt * PT))


def note(msg):
    print(msg)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cSld = s._element.find(qn('p:cSld'))
    bg = etree.SubElement(cSld, qn('p:bg'))
    cSld.remove(bg)
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, qn('p:bgPr'))
    fill = etree.SubElement(bgPr, qn('a:solidFill'))
    clr = etree.SubElement(fill, qn('a:srgbClr'))
    clr.set('val', 'FFFFFF')
    etree.SubElement(bgPr, qn('a:effectLst'))
    return s


def add_text(slide, text, left_pt, baseline_pt, size, color, bold=False,
             align=PP_ALIGN.LEFT, width_pt=None, font=FONT):
    """Top-anchored zero-inset text box whose first basline sits at baseline_pt."""
    w = width_pt if width_pt is not None else (metrics.advance(text, size, bold) + 6)
    top = baseline_pt - size
    tb = slide.shapes.add_textbox(Emu(emu(left_pt)), Emu(emu(top)), Emu(emu(w)),
                                  Emu(emu(size * 2)))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        rPr.append(rPr.makeelement(qn(tag), {"typeface": font}))
    return tb


def add_svg_pic(slide, svg_path, png_path, left_pt, top_pt, w_pt, h_pt):
    """Picture whose primary format is vector SVG, with a high-dpi PNG fallback."""
    pic = slide.shapes.add_picture(png_path, Emu(emu(left_pt)), Emu(emu(top_pt)),
                                   Emu(emu(w_pt)), Emu(emu(h_pt)))
    pkg = slide.part.package
    partname = pkg.next_partname(PackURI('/ppt/media/image%d.svg'))
    svg_part = Part(partname, 'image/svg+xml', pkg, open(svg_path, 'rb').read())
    rId = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = pic._element.blipFill.blip
    extLst = etree.SubElement(blip, qn('a:extLst'))
    ext = etree.SubElement(extLst, qn('a:ext'))
    ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    sb = etree.SubElement(ext, '{%s}svgBlip' % NS_ASVG)
    sb.set('{%s}embed' % RNS, rId)
    pic._element.nvPicPr.cNvPr.set('descr', 'LaTeX formula (vector SVG)')
    pic._element.nvPicPr.cNvPr.set('name', 'LaTeX formula (vector SVG)')
    return pic


def assert_uniform(text, where):
    """Left-aligned runs must be pure CJK or pure ASCII: a mixed run makes the
    renderer inject East-Asian/Latin auto-spacing, which no longer matches the
    measured advance width."""
    has_ascii = any(ord(c) < 128 for c in text)
    has_wide = any(ord(c) >= 128 for c in text)
    if has_ascii and has_wide:
        raise AssertionError("mixed-width run %r on %s would shift layout" % (text, where))


def polish(path, title, nslides):
    """Strip python-pptx's default-template metadata from the saved package.

    The stock template ships a blank 4:3 thumbnail, "On-screen Show (4:3)",
    <Slides>0</Slides> and a third-party author string.  None of that describes
    this deck, and a wrong embedded thumbnail is what Windows Explorer shows in
    the file preview, so the stale one is dropped (Explorer/PowerPoint then
    generate their own).
    """
    import re, shutil, zipfile
    tmp = path + ".tmp"
    zin = zipfile.ZipFile(path)
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        name = item.filename
        if name == "docProps/thumbnail.jpeg":
            continue
        data = zin.read(name)
        if name == "docProps/core.xml":
            s = data.decode("utf-8")
            s = re.sub(r"<dc:title/>", "<dc:title>%s</dc:title>" % title, s)
            s = re.sub(r"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>",
                       "<cp:lastModifiedBy></cp:lastModifiedBy>", s)
            s = re.sub(r"<dc:description>[^<]*</dc:description>",
                       "<dc:description></dc:description>", s)
            data = s.encode("utf-8")
        elif name == "docProps/app.xml":
            s = data.decode("utf-8")
            s = re.sub(r"<PresentationFormat>[^<]*</PresentationFormat>",
                       "<PresentationFormat>On-screen Show (16:9)</PresentationFormat>", s)
            s = re.sub(r"<Slides>\d+</Slides>", "<Slides>%d</Slides>" % nslides, s)
            s = re.sub(r"<Application>[^<]*</Application>",
                       "<Application>Microsoft Office PowerPoint</Application>", s)
            s = re.sub(r"<TotalTime>\d+</TotalTime>", "<TotalTime>0</TotalTime>", s)
            data = s.encode("utf-8")
        elif name == "_rels/.rels":
            s = data.decode("utf-8")
            s = re.sub(r'<Relationship[^>]*Type="[^"]*thumbnail"[^>]*/>', "", s)
            data = s.encode("utf-8")
        zout.writestr(item, data)
    zout.close()
    zin.close()
    shutil.move(tmp, path)


def main():
    C, L, man, plans = layout.build()
    ok, problems = layout.check(C, L, man, plans)
    for p in problems:
        note("  LAYOUT PROBLEM: " + p)
    if not ok:
        raise SystemExit("layout checks failed; fix content.json before building")

    prs = Presentation()
    prs.slide_width = Emu(emu(L["slide_w"]))
    prs.slide_height = Emu(emu(L["slide_h"]))

    # ---------------- cover ----------------
    cv = C["cover"]
    s = new_slide(prs)
    add_text(s, cv["title"], 0, cv["title_baseline"], cv["title_size"], cv["title_color"],
             bold=True, align=PP_ALIGN.CENTER, width_pt=L["slide_w"])
    add_text(s, cv["subtitle"], 0, cv["subtitle_baseline"], cv["subtitle_size"],
             cv["subtitle_color"], bold=False, align=PP_ALIGN.CENTER, width_pt=L["slide_w"])

    inventory = [{"n": 1, "title": cv["title"], "kind": "cover",
                  "content_bottom": round(cv["subtitle_baseline"] + 0.12 * cv["subtitle_size"], 2),
                  "texts": [cv["title"], cv["subtitle"]], "math": []}]

    # ---------------- question slides ----------------
    for sl, p in zip(C["slides"], plans):
        s = new_slide(prs)
        add_text(s, p["title"], L["margin_l"], L["title_baseline"], L["title_size"],
                 L["title_color"], bold=True)
        txts = []
        for t in p["texts"]:
            assert_uniform(t["s"], "slide %d" % p["n"])
            add_text(s, t["s"], t["x"], t["base"], t["size"],
                     L["label_color"] if t["bold"] else L["stem_color"], bold=t["bold"])
            txts.append(t["s"])
        for m in p["math"]:
            md = man[m["name"]]
            add_svg_pic(s, os.path.join(HERE, md["svg"]), os.path.join(HERE, md["png"]),
                        m["x"], m["base"] - md["ht_disp"], md["wd_disp"],
                        md["ht_disp"] + md["dp_disp"])
        inventory.append({"n": p["n"], "title": p["title"], "kind": "question",
                          "content_bottom": p["content_bottom"], "texts": txts,
                          "math": [m["name"] for m in p["math"]]})

    out = os.path.join(OUT, OUTNAME)
    prs.save(out)
    polish(out, C["cover"]["title"], len(prs.slides._sldIdLst))
    json.dump(inventory, open(os.path.join(HERE, "inventory.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    nslides = len(prs.slides._sldIdLst)
    print("saved: %s" % out)
    print("size : %.1f KB, slides = %d" % (os.path.getsize(out) / 1024.0, nslides))
    for p in plans:
        print("  slide %2d  %-24s contentBottom=%6.1fpt  blankArea=%5.1f%%"
              % (p["n"], p["title"], p["content_bottom"],
                 (L["slide_h"] - p["content_bottom"]) / L["slide_h"] * 100))
    if nslides != 10:
        raise SystemExit("expected exactly 10 slides, got %d" % nslides)
    return 0


if __name__ == "__main__":
    sys.exit(main())
