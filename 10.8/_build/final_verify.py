# -*- coding: utf-8 -*-
import os, sys, json, zipfile
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
BASE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.dirname(BASE)
PPTX = os.path.join(OUT, "\u590d\u6570\u4e13\u9898_\u8bfe\u5802\u677f\u4e66\u7248.pptx")
prs = Presentation(PPTX)
BAD = ["\u7b54\u6848","\u89e3\u6790","\u601d\u8def","\u63d0\u793a","\u89e3\uff1a","\u7efc\u4e0a","\u6240\u4ee5","\u7ec3\u4e60","\u6c34\u5370","\u9875\u7801"]
issues = []
print("slide size: %.1f x %.1f pt  ratio=%.4f  slides=%d" %
      (prs.slide_width/12700, prs.slide_height/12700, prs.slide_width/prs.slide_height, len(prs.slides._sldIdLst)))
for i, s in enumerate(prs.slides, 1):
    txts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    npic = sum(1 for sh in s.shapes if sh.__class__.__name__ == "Picture")
    nsh  = len(s.shapes._spTree) - 2
    print("slide %d: shapes=%2d pictures=%d  ::  %s" % (i, nsh, npic, " / ".join(txts)))
    j = " ".join(txts)
    for b in BAD:
        if b in j: issues.append((i, b))
    if i == 1:
        if "".join(txts) != "\u590d\u6570": issues.append((1, "cover != 复数"))
    else:
        if j.count("\u4f8b\u9898") != 1: issues.append((i, "例题 count"))
        first = sorted([(sh.top, sh.text_frame.text) for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()])[0][1]
        if not first.startswith("\u9898\u578b"): issues.append((i, "title first"))
        if npic < 1: issues.append((i, "no formula image"))
with zipfile.ZipFile(PPTX) as z:
    svg = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".svg")]
    png = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".png")]
print("media: svg=%d (one per formula picture)  png=%d (deduplicated fallbacks, identical bytes render identically)" % (len(svg), len(png)))
print("ISSUES:", issues if issues else "NONE")
