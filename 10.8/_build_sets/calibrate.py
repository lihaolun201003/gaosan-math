# -*- coding: utf-8 -*-
"""Empirical calibration of (font, size) -> baseline offset from textbox top,
plus advance-width validation for the CJK / Latin glyphs used by the deck.

Builds a probe deck, renders it with the bundled LibreOffice kit, and measures
the ink bands so the deck builder can place math baselines exactly on a text
baseline instead of guessing font ascent metrics.
"""
import json
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.util import Pt
from PIL import Image

from deckkit import add_textbox, patch_theme_east_asian

BASE = os.path.dirname(os.path.abspath(__file__))
NODE = r"C:\Users\lihao\AppData\Local\Programs\DeepSeek Harness\resources\runtime\primary-runtime\dependencies\node\bin\node.exe"
CLI = r"C:\Users\lihao\AppData\Local\Programs\DeepSeek Harness\resources\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit\lib\cli.js"

SLIDE_W, SLIDE_H = 960.0, 540.0
DPI = 300.0
K = DPI / 72.0                       # px per pt

SIMSUN = (r"C:\Windows\Fonts\simsun.ttc", 0)
MSYHBD = (r"C:\Windows\Fonts\msyhbd.ttc", 0)

# label, font file, face name, size, textbox top
CASES = [
    ("sun24", SIMSUN, "宋体", 24.0, 30.0),
    ("sun26", SIMSUN, "宋体", 26.0, 120.0),
    ("sun28", SIMSUN, "宋体", 28.0, 210.0),
    ("yahei32", MSYHBD, "微软雅黑", 32.0, 300.0),
    ("yahei30", MSYHBD, "微软雅黑", 30.0, 390.0),
    ("sun26b", SIMSUN, "宋体", 26.0, 470.0),
]

BOXES = [("cjk1", "已", 60.0), ("lat1", "x", 200.0),
         ("cjk10", "已" * 10, 380.0), ("lat10", "x" * 10, 760.0)]


def build_probe(path):
    prs = Presentation()
    prs.slide_width = Pt(SLIDE_W)
    prs.slide_height = Pt(SLIDE_H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for label, (ffile, fidx), face, size, top in CASES:
        for _bname, text, left in BOXES:
            add_textbox(s, left, top, 300.0, size * 2.0, text, face, size,
                        color="000000")
    prs.save(path)
    patch_theme_east_asian(path, "Times New Roman", "宋体")
    return path


def render(pptx, outdir, dpi=DPI):
    if os.path.isdir(outdir):
        import shutil
        shutil.rmtree(outdir)
    cmd = [NODE, CLI, "render", "--input", pptx, "--output-dir", outdir,
           "--dpi", str(int(dpi)), "--pages", "1"]
    p = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    out = p.stdout.decode("utf-8", "replace")
    print(out[-1200:])
    if p.returncode != 0:
        raise RuntimeError("render failed")
    files = [f for f in os.listdir(outdir) if f.lower().endswith(".png")]
    return os.path.join(outdir, files[0])


def bands(img, y0, y1, x0=0, x1=None, gap=3):
    """Column bands of ink inside a y-window, merged across gaps < `gap` px."""
    W, H = img.size
    x1 = W if x1 is None else x1
    y0 = max(0, int(y0)); y1 = min(H, int(y1))
    px = img.load()
    cols = []
    for x in range(int(x0), int(x1)):
        hit = False
        for y in range(y0, y1):
            if px[x, y] < 128:
                hit = True
                break
        cols.append(hit)
    out, i = [], 0
    while i < len(cols):
        if cols[i]:
            j = i
            while j < len(cols) and cols[j]:
                j += 1
            out.append((i + int(x0), j - 1 + int(x0)))
            i = j
        else:
            i += 1
    merged = []
    for b in out:                       # merge bands separated by < gap px
        if merged and b[0] - merged[-1][1] <= gap:
            merged[-1] = (merged[-1][0], b[1])
        else:
            merged.append(b)
    return merged


def ink_rows(img, x0, x1, y0, y1):
    px = img.load()
    rows = [y for y in range(int(y0), int(y1))
            if any(px[x, y] < 128 for x in range(int(x0), int(x1) + 1))]
    if not rows:
        return None
    return rows[0], rows[-1]


def main():
    if len(sys.argv) > 1 and sys.argv[1] == "build":
        probe = os.path.join(BASE, "calib_probe.pptx")
        build_probe(probe)
        print("built", probe)
        return
    if len(sys.argv) > 2:
        img_path = sys.argv[2]
    else:
        img_path = os.path.join(BASE, "calib_out", "page-0001.png")
    im = Image.open(img_path).convert("L")
    W, H = im.size
    print("render %s  %dx%d px  (%.2f px/pt)  slide %.0fx%.0f pt -> %.1fx%.1f pt"
          % (os.path.basename(img_path), W, H, K, SLIDE_W, SLIDE_H, W / K, H / K))

    result = {}
    for label, (ffile, fidx), face, size, top in CASES:
        y0, y1 = (top - 12) * K, (top + size * 2.1) * K
        bs = bands(im, y0, y1, gap=0.5 * size * K)
        if len(bs) != 4:
            print("  %-9s expected 4 bands, got %d -> %s" % (label, len(bs), bs))
            continue
        rec = {"face": face, "size": size, "top": top}
        # baseline = ink bottom of the single latin 'x'
        b0, b1 = bs[1]
        rows = ink_rows(im, b0, b1, y0, y1)
        baseline = (rows[1] + 1) / K
        rec["baseline_px"] = rows[1] + 1
        rec["baseline_offset"] = baseline - top
        rec["baseline_em"] = (baseline - top) / size
        # CJK ink box relative to baseline
        r = ink_rows(im, bs[0][0], bs[0][1], y0, y1)
        rec["cjk_top_from_baseline"] = (r[0] / K) - baseline
        rec["cjk_bot_from_baseline"] = (r[1] + 1) / K - baseline
        rec["cjk_ink_h"] = (r[1] - r[0] + 1) / K
        # advances
        w_cjk10 = (bs[2][1] - bs[2][0] + 1) / K
        w_cjk1 = (bs[0][1] - bs[0][0] + 1) / K
        rec["adv_cjk"] = (w_cjk10 - w_cjk1) / 9.0
        w_lat10 = (bs[3][1] - bs[3][0] + 1) / K
        w_lat1 = (bs[1][1] - bs[1][0] + 1) / K
        rec["adv_lat"] = (w_lat10 - w_lat1) / 9.0
        rec["band_x"] = [(b[0] / K, b[1] / K) for b in bs]
        result[label] = rec
        print("  %-9s %-6s %4.1fpt  baselineFromTop=%6.3fpt (%.4f em)  "
              "adv_cjk=%6.3f adv_x=%6.3f  cjk_ink_h=%5.2f"
              % (label, face, size, rec["baseline_offset"], rec["baseline_em"],
                 rec["adv_cjk"], rec["adv_lat"], rec["cjk_ink_h"]))

    # --- cross-check against PIL-computed advances -------------------------
    from PIL import ImageFont
    print("\nPIL cross-check (px at 1000, scaled to pt):")
    for label, (ffile, fidx), face, size, top in CASES:
        if label not in result:
            continue
        f = ImageFont.truetype(ffile, 1000, index=fidx)
        adv_cjk = f.getlength("已") / 1000.0 * size
        adv_x = f.getlength("x") / 1000.0 * size
        adv_kai = f.getlength("已知集合 U={1,2,3,4}，ACD．（）a") / 1000.0 * size
        print("  %-9s PIL adv_cjk=%6.3f (render %6.3f, d=%+.3f)   "
              "PIL adv_x=%6.3f (render %6.3f, d=%+.3f)"
              % (label, adv_cjk, result[label]["adv_cjk"],
                 adv_cjk - result[label]["adv_cjk"], adv_x,
                 result[label]["adv_lat"], adv_x - result[label]["adv_lat"]))

    with open(os.path.join(BASE, "calib_text.json"), "w", encoding="utf-8") as fh:
        json.dump({"dpi": DPI, "cases": result}, fh, ensure_ascii=False, indent=1)
    print("\nwrote calib_text.json")


if __name__ == "__main__":
    main()
