# -*- coding: utf-8 -*-
"""Shared PPTX helpers: CJK-safe run fonts, exact textboxes, svgBlip pictures."""
import copy
import os

from pptx.util import Pt, Emu
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from lxml import etree

PT = 12700                      # EMU per point
EMU_PER_PT = 12700

NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def emu(pt):
    return Emu(int(round(pt * EMU_PER_PT)))


def set_run_font(run, name, size=None, bold=None, italic=None, color=None,
                 spacing=None):
    """Set latin + east-asian + complex-script typefaces on one run."""
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)
    f.name = name                                  # writes a:latin
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.insert_element_before(el, "a:sym", "a:hlinkClick",
                                      "a:hlinkMouseOver", "a:rtl", "a:extLst")
        el.set("typeface", name)
    if spacing is not None:                        # character spacing, in 1/100 pt
        rPr.set("spc", str(int(round(spacing * 100))))


def add_textbox(slide, left_pt, top_pt, width_pt, height_pt, text, font,
                size, bold=False, color="000000", align=None, anchor=None,
                wrap=False, line_spacing=None, space_after=0):
    """A textbox with zero insets, no autofit and (by default) no wrapping.

    Wrapping is off by default so the renderer never re-flows a measured
    segment; width is only a container hint.
    """
    box = slide.shapes.add_textbox(emu(left_pt), emu(top_pt), emu(width_pt), emu(height_pt))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor if anchor is not None else MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing is not None:
        p.line_spacing = line_spacing
    if space_after:
        p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, bold=bold, color=color)
    return box


def add_svg_picture(slide, svg_path, png_path, left_pt, top_pt, width_pt, height_pt,
                    alt=None):
    """Insert a picture whose primary blip is the raster fallback and which also
    carries an svgBlip extension, so Office renders the vector original.

    `alt` (usually the LaTeX source) is written into both the shape name and the
    alt text so the deck stays self-documenting when reviewed or edited.
    """
    pic = slide.shapes.add_picture(png_path, emu(left_pt), emu(top_pt),
                                   emu(width_pt), emu(height_pt))
    if alt:
        cNvPr = pic._element.nvPicPr.cNvPr
        cNvPr.set("name", alt)
        cNvPr.set("descr", alt)
    pkg = slide.part.package
    partname = pkg.next_partname(PackURI("/ppt/media/image%d.svg"))
    with open(svg_path, "rb") as fh:
        blob = fh.read()
    svg_part = Part(partname, "image/svg+xml", pkg, blob)
    rId = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = pic._element.blipFill.blip
    ext_lst = blip.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(blip, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", SVG_EXT_URI)
    svg_blip = etree.SubElement(ext, "{%s}svgBlip" % NS_ASVG)
    svg_blip.set("{%s}embed" % REL_NS, rId)
    return pic


def finalize_pptx(pptx_path, latin_font, ea_font, slide_count, title, author,
                  subject):
    """Rewrite package-level details in an already-saved .pptx:

    * theme latin / east-asian typefaces (so runs lacking a:ea still show CJK
      in the intended face),
    * `p:sldSz/@type` and the `PresentationFormat` label, which python-pptx
      leaves at its 4:3 template default even after the slide size is changed,
    * `Slides` count and the template's placeholder author / description,
    * the stale 4:3 blank thumbnail shipped by the default template.
    """
    import re
    import shutil
    import zipfile

    tmp = pptx_path + ".tmp"
    with zipfile.ZipFile(pptx_path, "r") as src:
        names = src.namelist()
        data = {n: src.read(n) for n in names}
        infos = {i.filename: i for i in src.infolist()}

    # --- theme fonts -------------------------------------------------------
    themes = [n for n in names if re.match(r"ppt/theme/theme\d+\.xml$", n)]
    for name in themes:
        xml = data[name].decode("utf-8")
        for scheme in ("majorFont", "minorFont"):
            m = re.search(r"<a:%s>(.*?)</a:%s>" % (scheme, scheme), xml, re.S)
            if not m:
                continue
            block = m.group(0)
            new_block = re.sub(r"<a:latin typeface=\"[^\"]*\"",
                               '<a:latin typeface="%s"' % latin_font, block, count=1)
            if "<a:ea " in new_block:
                new_block = re.sub(r"<a:ea typeface=\"[^\"]*\"",
                                   '<a:ea typeface="%s"' % ea_font, new_block, count=1)
            else:
                new_block = new_block.replace(
                    "<a:cs ", '<a:ea typeface="%s"/><a:cs ' % ea_font, 1)
                if "<a:ea " not in new_block:
                    new_block = re.sub(r"(<a:latin typeface=\"[^\"]*\"/>)",
                                       r'\1<a:ea typeface="%s"/>' % ea_font,
                                       new_block, count=1)
            xml = xml.replace(block, new_block)
        data[name] = xml.encode("utf-8")

    # --- slide size label --------------------------------------------------
    pres = "ppt/presentation.xml"
    if pres in data:
        xml = data[pres].decode("utf-8")
        xml = re.sub(r'(<p:sldSz\b[^>]*?)\stype="[^"]*"', r"\1", xml)
        xml = xml.replace("<p:sldSz", '<p:sldSz type="screen16x9"', 1)
        data[pres] = xml.encode("utf-8")

    # --- app / core properties --------------------------------------------
    app = "docProps/app.xml"
    if app in data:
        xml = data[app].decode("utf-8")
        xml = re.sub(r"<PresentationFormat>.*?</PresentationFormat>",
                     "<PresentationFormat>On-screen Show (16:9)</PresentationFormat>",
                     xml, flags=re.S)
        xml = re.sub(r"<Slides>\d+</Slides>", "<Slides>%d</Slides>" % slide_count, xml)
        xml = re.sub(r"<Application>.*?</Application>",
                     "<Application>Microsoft Office PowerPoint</Application>",
                     xml, flags=re.S)
        xml = re.sub(r"<Company>.*?</Company>", "<Company></Company>", xml, flags=re.S)
        xml = re.sub(r"<TotalTime>\d+</TotalTime>", "<TotalTime>0</TotalTime>", xml)
        data[app] = xml.encode("utf-8")

    core = "docProps/core.xml"
    if core in data:
        xml = data[core].decode("utf-8")
        xml = re.sub(r"<dc:title\s*/>", "<dc:title>%s</dc:title>" % title, xml)
        xml = re.sub(r"<dc:title>.*?</dc:title>", "<dc:title>%s</dc:title>" % title,
                     xml, flags=re.S)
        if "<dc:creator" in xml:
            xml = re.sub(r"<dc:creator\s*/>", "<dc:creator>%s</dc:creator>" % author, xml)
            xml = re.sub(r"<dc:creator>.*?</dc:creator>",
                         "<dc:creator>%s</dc:creator>" % author, xml, flags=re.S)
        else:
            xml = xml.replace("<dc:title>", "<dc:creator>%s</dc:creator><dc:title>" % author, 1)
        xml = re.sub(r"<cp:lastModifiedBy\s*/>",
                     "<cp:lastModifiedBy>%s</cp:lastModifiedBy>" % author, xml)
        xml = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                     "<cp:lastModifiedBy>%s</cp:lastModifiedBy>" % author,
                     xml, flags=re.S)
        if "<dc:description" in xml:
            xml = re.sub(r"<dc:description\s*/>",
                         "<dc:description>%s</dc:description>" % subject, xml)
            xml = re.sub(r"<dc:description>.*?</dc:description>",
                         "<dc:description>%s</dc:description>" % subject,
                         xml, flags=re.S)
        else:
            xml = xml.replace(
                "<cp:lastModifiedBy>",
                "<dc:description>%s</dc:description><cp:lastModifiedBy>" % subject, 1)
        data[core] = xml.encode("utf-8")

    # --- drop the template's blank 4:3 thumbnail ---------------------------
    drop = {"docProps/thumbnail.jpeg", "docProps/thumbnail.jpg"}
    rels = "_rels/.rels"
    if rels in data:
        xml = data[rels].decode("utf-8")
        xml = re.sub(r"<Relationship\b[^>]*?/thumbnail\"[^>]*?/>", "", xml)
        data[rels] = xml.encode("utf-8")

    names = [n for n in names if n not in drop]

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for name in names:
            info = infos[name]
            zi = zipfile.ZipInfo(name, date_time=info.date_time)
            zi.compress_type = info.compress_type
            zi.external_attr = info.external_attr
            out.writestr(zi, data[name])
    shutil.move(tmp, pptx_path)
    return {"themes": themes, "dropped": sorted(drop)}


